"""
Excel Values-Only Converter
Upload an Excel file that was recalculated and saved in Excel.
Every formula is replaced by its last calculated value, external links are removed,
and formatting is kept. Download the light copy.

Run:  streamlit run app.py
"""
import io
from datetime import datetime

import streamlit as st
from openpyxl import load_workbook


def to_values_only(file_bytes: bytes, is_xlsm: bool):
    """Return (new_file_bytes, report dict)."""
    wb = load_workbook(io.BytesIO(file_bytes), keep_vba=is_xlsm)            # formulas + formatting
    wb_vals = load_workbook(io.BytesIO(file_bytes), data_only=True)         # cached values

    report = {"formulas_replaced": 0, "missing_values": [], "sheets": []}

    for ws in wb.worksheets:
        ws_vals = wb_vals[ws.title]
        count = 0
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f":                    # formula (incl. array formulas)
                    value = ws_vals[cell.coordinate].value
                    if value is None and len(report["missing_values"]) < 50:
                        report["missing_values"].append(f"{ws.title}!{cell.coordinate}")
                    cell.value = value
                    count += 1
        report["formulas_replaced"] += count
        report["sheets"].append((ws.title, count))

    # Remove external workbook links so the copy never asks to "update links"
    report["external_links_removed"] = len(getattr(wb, "_external_links", []))
    wb._external_links = []

    # Drop named ranges that still point to other workbooks
    removed_names = []
    for name in list(wb.defined_names.keys()):
        if "[" in str(wb.defined_names[name].attr_text):
            del wb.defined_names[name]
            removed_names.append(name)
    report["external_names_removed"] = removed_names

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), report


# ---------------- UI ----------------
st.set_page_config(page_title="MRP Values-Only Converter", page_icon="📦")
st.title("📦 MRP Values-Only Converter")
st.write(
    "Upload the MRP file **after** it has been updated and saved in Excel. "
    "You'll get a light copy with every formula replaced by its current value."
)

uploaded = st.file_uploader("MRP file", type=["xlsx", "xlsm"])

if uploaded:
    is_xlsm = uploaded.name.lower().endswith(".xlsm")
    with st.spinner("Converting… large files can take a minute"):
        new_bytes, rep = to_values_only(uploaded.getvalue(), is_xlsm)

    st.success(f"Done: {rep['formulas_replaced']:,} formulas replaced with values.")

    c1, c2, c3 = st.columns(3)
    c1.metric("Original size", f"{len(uploaded.getvalue())/1e6:.1f} MB")
    c2.metric("New size", f"{len(new_bytes)/1e6:.1f} MB")
    c3.metric("External links removed", rep["external_links_removed"])

    with st.expander("Formulas replaced per sheet"):
        st.table({"Sheet": [s for s, _ in rep["sheets"]],
                  "Formulas": [n for _, n in rep["sheets"]]})

    if rep["missing_values"]:
        st.warning(
            "Some formula cells had no saved value (shown first 50). This usually means the "
            "file wasn't recalculated and saved in Excel, or those formulas return blank. "
            "Open it in Excel, press Ctrl+Alt+F9, save, and upload again if they matter.\n\n"
            + ", ".join(rep["missing_values"])
        )

    stem = uploaded.name.rsplit(".", 1)[0]
    ext = "xlsm" if is_xlsm else "xlsx"
    st.download_button(
        "⬇️ Download values-only file",
        data=new_bytes,
        file_name=f"{stem}_VALUES_{datetime.now():%Y-%m-%d}.{ext}",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

st.caption(
    "Note: charts, images, slicers and pivot caches may not carry over. Cell values, "
    "formatting, column widths and conditional formatting do."
)
